"""Build the competition Word document from its reviewed Markdown source.

Use the Codex bundled Python (python-docx/Pillow) and ffmpeg on PATH.
Generated DOCX, extracted video frames and QA renders remain untracked.
"""
from pathlib import Path
import argparse
import re
import shutil
import subprocess

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / '资料包/02-作品说明文档/参赛说明.md'
OUTPUT = ROOT / 'deliverables/competition-doc/菌灯之下_游戏参赛说明.docx'


def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def table(doc, rows):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [4.0, 12.6]
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement('w:' + edge)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            node.set(qn('w:' + k), v)
        borders.append(node)
    pr.append(borders)
    for i, values in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = Pt(14)
            for r in p.runs:
                r.font.size = Pt(10)
                r.bold = i == 0
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DFEAF0' if i == 0 else ('F5F7F8' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                m = OxmlElement('w:' + edge)
                m.set(qn('w:w'), '90')
                m.set(qn('w:type'), 'dxa')
                margins.append(m)
            tcpr.append(margins)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.add_run().font.size = Pt(1)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    frames = ROOT / 'qa/local/competition-doc'
    frames.mkdir(parents=True, exist_ok=True)
    ffmpeg = shutil.which('ffmpeg')
    for second in [6, 30, 42, 66]:
        target = frames / f'video-{second:03}.png'
        if not target.exists():
            if not ffmpeg:
                raise RuntimeError('ffmpeg is required to extract original gameplay frames')
            subprocess.run([ffmpeg, '-hide_banner', '-loglevel', 'error', '-ss', str(second),
                            '-i', str(ROOT / '资料包/FungusLight-Gameplay-演示视频.mp4'),
                            '-frames:v', '1', '-y', str(target)], check=True)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.65)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.footer_distance = Cm(0.75)
    font(doc.styles['Normal'], 10.5)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = Pt(16)
    normal.space_after = Pt(7)
    normal.widow_control = True
    for name, size in [('Title', 29), ('Subtitle', 16), ('Heading 1', 19), ('Heading 2', 12.5)]:
        font(doc.styles[name], size, name != 'Subtitle')
        doc.styles[name].font.italic = False
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(10), Pt(8)
        pf.keep_with_next = True
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for grid in sec._sectPr.findall(qn('w:docGrid')):
        sec._sectPr.remove(grid)
    font(doc.styles['Caption'], 8.5)
    doc.styles['Caption'].font.color.rgb = RGBColor.from_string('555555')
    doc.styles['Caption'].paragraph_format.space_after = Pt(10)
    doc.styles['Caption'].paragraph_format.line_spacing = 1.15

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('菌灯之下  |  ')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('777777')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = '菌灯之下 游戏参赛说明与 AI 创作记录'
    doc.core_properties.subject = '游戏背景 玩法特色 AI 创作流程与验证'
    doc.core_properties.author = 'FungusLight 项目'
    doc.core_properties.keywords = '菌灯之下, AI游戏创作, 动作RPG, Godot'

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    first = True
    page_pending = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == '<!-- PAGE -->':
            page_pending = True
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            p = doc.add_paragraph(line[3:], 'Subtitle' if first else 'Heading 1')
            p.paragraph_format.page_break_before = page_pending
            page_pending = False
            first = False
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('|'):
            rows = [[x.strip() for x in line.strip('|').split('|')]]
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                i += 1
                if all(re.fullmatch(r':?-+:?', x) for x in row):
                    continue
                rows.append(row)
            table(doc, rows)
        elif line.startswith('!['):
            m = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
            path = (SOURCE.parent / m.group(2)).resolve()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            pic = p.add_run().add_picture(str(path), width=Cm(16.6))
            pic._inline.docPr.set('descr', m.group(1))
        elif line.startswith('图 '):
            p = doc.add_paragraph(line, 'Caption')
            p.paragraph_format.keep_together = True
        else:
            p = doc.add_paragraph(line)
            if re.match(r'^\d+\. ', line):
                p.paragraph_format.space_after = Pt(5)
    for p in doc.element.xpath('.//w:p'):
        ppr = p.get_or_add_pPr()
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'), '0')
        ppr.append(snap)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == '__main__':
    main()
